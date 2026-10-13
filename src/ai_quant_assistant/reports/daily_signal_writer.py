from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pandas as pd
from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl.styles import Font, PatternFill

from ai_quant_assistant.reports.presentation import (
    TERM_CATALOG,
    install_filters,
    rule_label,
)

FRAME_COLUMNS = {
    "signals": [
        "signal_id",
        "signal_date",
        "planned_execution_date",
        "symbol",
        "action",
        "status",
        "signal_price",
        "candidate_eligible",
        "candidate_score",
        "candidate_rank",
        "planned_shares",
        "rule_evidence",
    ],
    "orders": [
        "signal_id",
        "signal_date",
        "planned_execution_date",
        "execution_date",
        "symbol",
        "side",
        "status",
        "shares",
        "signal_price",
        "execution_price",
        "notional",
        "cost",
        "reason",
        "next_action",
        "candidate_score",
        "candidate_rank",
        "planned_shares",
        "cash_before",
        "cash_after",
        "position_shares_before",
        "position_shares_after",
        "commission_cost",
        "slippage_cost",
        "stamp_duty_cost",
        "retry_continues",
    ],
    "positions": [
        "date",
        "symbol",
        "shares",
        "price",
        "market_value",
        "candidate_score",
        "candidate_rank",
        "pending_exit",
    ],
    "candidates": [
        "signal_date",
        "symbol",
        "composite_score",
        "rank",
        "selected",
        "component_values",
        "component_sources",
    ],
}


def _value(value: Any) -> Any:
    if pd.isna(value):
        return None
    return value.item() if hasattr(value, "item") else value


def _records(frame: pd.DataFrame) -> list[dict[str, Any]]:
    return [
        {key: _value(value) for key, value in row.items()}
        for row in frame.to_dict(orient="records")
    ]


def daily_signal_result_view(
    *,
    spec: dict[str, Any],
    metrics: dict[str, Any],
    nav: pd.DataFrame,
    orders: pd.DataFrame,
    positions: pd.DataFrame,
    diagnostics: dict[str, Any],
) -> dict[str, Any]:
    """Build a display-only projection from saved outputs without changing calculations."""
    final_nav = nav.iloc[-1].to_dict() if not nav.empty else {}
    executed = orders.loc[orders.get("status", pd.Series(dtype=object)).eq("executed")]
    unfilled = orders.loc[~orders.index.isin(executed.index)]
    final_positions = pd.DataFrame(columns=positions.columns)
    if not positions.empty and "date" in positions:
        final_date = positions["date"].dropna().max()
        if final_date is not None:
            final_positions = positions.loc[positions["date"].eq(final_date)].copy()
    final_value = _value(final_nav.get("portfolio_value", metrics.get("final_value")))
    position_rows = []
    for row in _records(final_positions):
        market_value = _value(row.get("market_value"))
        ratio = None
        if (
            isinstance(final_value, (int, float))
            and final_value > 0
            and isinstance(market_value, (int, float))
        ):
            ratio = market_value / final_value
        position_rows.append(
            {
                **{key: _value(value) for key, value in row.items()},
                "asset_ratio": ratio,
            }
        )
    open_positions = diagnostics.get("open_positions")
    open_exits = diagnostics.get("open_exit_intents")
    account = {
        "initial_capital": spec.get("initial_capital"),
        "final_value": final_value,
        "final_cash": _value(final_nav.get("cash")),
        "final_position_value": _value(final_nav.get("position_value")),
        "total_return": metrics.get("total_return"),
        "benchmark_total_return": metrics.get("benchmark_total_return"),
        "excess_return": metrics.get("excess_return"),
        "max_drawdown": metrics.get("max_drawdown"),
        "total_transaction_cost": metrics.get("total_transaction_cost"),
        "signal_count": metrics.get("signal_count"),
        "order_count": len(orders),
        "buy_execution_count": int(executed.get("side", pd.Series(dtype=object)).eq("buy").sum()),
        "sell_execution_count": int(executed.get("side", pd.Series(dtype=object)).eq("sell").sum()),
        "unfilled_count": len(unfilled),
        "ending_position_count": len(open_positions)
        if isinstance(open_positions, list)
        else (len(final_positions) if not positions.empty else None),
        "pending_exit_count": len(open_exits) if isinstance(open_exits, list) else None,
    }
    return {
        "account": account,
        "executed": _records(executed.tail(80)),
        "unfilled": _records(unfilled.tail(80)),
        "final_positions": position_rows,
        "daily_accounts": _records(nav.tail(80)),
    }


def daily_signal_trust_summary(
    *, spec: dict[str, Any], metrics: dict[str, Any], nav: pd.DataFrame,
    orders: pd.DataFrame, diagnostics: dict[str, Any], reproducibility: dict[str, Any],
) -> dict[str, Any]:
    """Persist facts from the saved input/output frames; never infer trade P&L."""
    trades = orders.loc[orders["status"].eq("executed")] if "status" in orders else orders.iloc[0:0]
    symbols = []
    if "symbol" in trades:
        for symbol, group in trades.groupby("symbol", sort=True):
            buys = group.loc[group["side"].eq("buy")]
            sells = group.loc[group["side"].eq("sell")]
            def total(frame: pd.DataFrame, column: str) -> float | None:
                if column not in frame or frame[column].isna().any():
                    return None
                return float(pd.to_numeric(frame[column], errors="coerce").sum())
            symbols.append({
                "symbol": str(symbol), "buy_count": len(buys), "sell_count": len(sells),
                "buy_notional": total(buys, "notional"),
                "sell_notional": total(sells, "notional"),
                "total_cost": total(group, "cost"),
            })
    drawdown = None
    if not nav.empty and {"date", "portfolio_nav"}.issubset(nav.columns):
        values = pd.to_numeric(nav["portfolio_nav"], errors="coerce")
        if values.notna().any():
            peaks = values.cummax()
            depth = values / peaks - 1
            trough = depth.idxmin()
            peak = values.loc[:trough].idxmax()
            columns = ("date", "cash", "position_value", "portfolio_value")
            drawdown = {
                "peak_date": str(nav.loc[peak, "date"])[:10],
                "trough_date": str(nav.loc[trough, "date"])[:10],
                "magnitude": float(-depth.loc[trough]),
                "peak_account": {key: (str(nav.loc[peak, key])[:10] if key == "date" else _value(nav.loc[peak, key])) if key in nav else None for key in columns},
                "trough_account": {key: (str(nav.loc[trough, key])[:10] if key == "date" else _value(nav.loc[trough, key])) if key in nav else None for key in columns},
            }
    capital = spec.get("initial_capital")
    cost = metrics.get("total_transaction_cost")
    final = metrics.get("final_value")
    snapshot = spec.get("preflight_snapshot") or {}
    return {
        "scope": {
            "pool_id": (spec.get("stock_pool") or {}).get("pool_id"),
            "pool_revision": (spec.get("stock_pool") or {}).get("pool_revision"),
            "frozen_count": len(spec.get("symbols") or []),
            "excluded": (spec.get("scope_adjustment") or {}).get("excluded"),
            "source_pool_id": (spec.get("scope_adjustment") or {}).get("pool_id"),
            "source_pool_revision": (spec.get("scope_adjustment") or {}).get("pool_revision"),
        },
        "research_start": spec.get("start_date"), "research_end": spec.get("end_date"),
        "effective_trading_days": len(nav) if not nav.empty else None,
        "warmup_start": reproducibility.get("technical_warmup_start"),
        "preflight": snapshot or None,
        "market_data_coverage": reproducibility.get("market_data_coverage"),
        "benchmark_data_coverage": reproducibility.get("benchmark_data_coverage"),
        "account_audit": diagnostics.get("account_audit"),
        "signal_count": metrics.get("signal_count"),
        "order_count": len(orders), "executed_count": len(trades),
        "unfilled_count": len(orders) - len(trades),
        "pending_exit_count": len(diagnostics["open_exit_intents"])
        if isinstance(diagnostics.get("open_exit_intents"), list) else None,
        "parameter_sha256": reproducibility.get("parameter_sha256"),
        "git_commit_sha": reproducibility.get("git_commit_sha"),
        "per_symbol": symbols,
        "drawdown_period": drawdown,
        "cost_to_initial_capital": float(cost) / float(capital)
        if isinstance(cost, (int, float)) and isinstance(capital, (int, float)) and capital > 0 else None,
        "cost_to_abs_net_gain": float(cost) / abs(float(final) - float(capital))
        if all(isinstance(value, (int, float)) for value in (cost, capital, final))
        and final != capital else None,
        "profit_attribution": None,
        "limitation": "静态或手工股票范围可能有幸存者偏差；未保存可核对的成本基础与配对规则，不计算单笔盈亏或股票收益贡献。",
    }


def daily_signal_chart(nav: pd.DataFrame, *, width: int = 900, height: int = 320) -> dict[str, Any]:
    """Prepare a bounded SVG model while preserving missing-value gaps."""
    if nav.empty:
        return {"width": width, "height": height, "segments": {}, "x_ticks": [], "y_ticks": []}
    frame = nav.copy()
    if len(frame) > 1200:
        indexes = sorted(
            {round(item * (len(frame) - 1) / 1199) for item in range(1200)} | {0, len(frame) - 1}
        )
        frame = frame.iloc[indexes]
    left, right, top, bottom = 74, 24, 28, 58
    values = pd.concat(
        [
            pd.to_numeric(
                frame[key] if key in frame else pd.Series(index=frame.index, dtype=float),
                errors="coerce",
            )
            for key in ("portfolio_nav", "benchmark_nav")
        ]
    ).dropna()
    if values.empty:
        return {"width": width, "height": height, "segments": {}, "x_ticks": [], "y_ticks": []}
    low, high = float(values.min()), float(values.max())
    padding = max((high - low) * 0.08, max(abs(low), abs(high), 1.0) * 0.01)
    low, high = low - padding, high + padding
    span = high - low or 1.0

    def x(index: int) -> float:
        return left + (width - left - right) * index / max(len(frame) - 1, 1)

    def y(value: float) -> float:
        return top + (height - top - bottom) * (high - value) / span

    segments: dict[str, list[str]] = {}
    for key in ("portfolio_nav", "benchmark_nav"):
        chunks, chunk = [], []
        for index, value in enumerate(pd.to_numeric(frame.get(key), errors="coerce")):
            if pd.notna(value):
                chunk.append(f"{x(index):.2f},{y(float(value)):.2f}")
            elif chunk:
                chunks.append(" ".join(chunk))
                chunk = []
        if chunk:
            chunks.append(" ".join(chunk))
        segments[key] = chunks
    x_count = min(6, len(frame))
    x_indexes = sorted(
        {round(item * (len(frame) - 1) / max(x_count - 1, 1)) for item in range(x_count)}
    )
    dates = frame.get("date", pd.Series(range(len(frame)))).astype(str).tolist()
    return {
        "width": width,
        "height": height,
        "plot": {"left": left, "right": width - right, "top": top, "bottom": height - bottom},
        "segments": segments,
        "x_ticks": [{"x": x(index), "label": dates[index]} for index in x_indexes],
        "y_ticks": [
            {"y": y(low + span * item / 5), "label": f"{low + span * item / 5:.3f}"}
            for item in range(6)
        ],
    }


def daily_signal_price_chart(
    indicators: pd.DataFrame,
    signals: pd.DataFrame,
    orders: pd.DataFrame,
    *,
    symbol: str | None = None,
    width: int = 900,
    height: int = 380,
) -> dict[str, Any]:
    """Build one bounded price-axis chart; execution prices remain event metadata only."""
    if indicators.empty or "symbol" not in indicators:
        return {"width": width, "height": height, "symbol": symbol, "segments": {}}
    symbols = sorted(str(value) for value in indicators["symbol"].dropna().unique())
    symbol = symbol if symbol in symbols else (symbols[0] if symbols else None)
    if symbol is None:
        return {"width": width, "height": height, "symbol": None, "segments": {}}
    frame = indicators.loc[indicators["symbol"].astype(str).eq(symbol)].copy()
    frame["date_key"] = frame["date"].astype(str).str[:10]
    related_signals = (
        signals.loc[signals.get("symbol", pd.Series(dtype=str)).astype(str).eq(symbol)]
        if not signals.empty
        else signals
    )
    related_orders = (
        orders.loc[orders.get("symbol", pd.Series(dtype=str)).astype(str).eq(symbol)]
        if not orders.empty
        else orders
    )
    key_dates = {
        str(row.get(key))[:10]
        for row in [*_records(related_signals), *_records(related_orders)]
        for key in ("signal_date", "planned_execution_date", "execution_date")
        if row.get(key) is not None
    }
    if len(frame) > 1_200:
        required = set(frame.index[frame["date_key"].isin(key_dates)]) | {
            frame.index[0],
            frame.index[-1],
        }
        remaining = max(0, 1_200 - len(required))
        sampled = [
            frame.index[round(index * (len(frame) - 1) / max(remaining - 1, 1))]
            for index in range(remaining)
        ]
        frame = frame.loc[sorted(required | set(sampled))]
    frame = frame.reset_index(drop=True)
    line_keys = [key for key in frame.columns if key == "signal_close" or key.startswith("ma_")]
    values = pd.concat([pd.to_numeric(frame[key], errors="coerce") for key in line_keys]).dropna()
    if values.empty:
        return {"width": width, "height": height, "symbol": symbol, "segments": {}}
    left, right, top, bottom = 74, 24, 28, 88
    low, high = float(values.min()), float(values.max())
    padding = max((high - low) * 0.08, max(abs(low), abs(high), 1.0) * 0.01)
    low, high = low - padding, high + padding
    span = high - low or 1.0

    def x(index: int) -> float:
        return left + (width - left - right) * index / max(len(frame) - 1, 1)

    def y(value: float) -> float:
        return top + (height - top - bottom) * (high - value) / span

    segments = {}
    for key in line_keys:
        chunks, chunk = [], []
        for index, value in enumerate(pd.to_numeric(frame[key], errors="coerce")):
            if pd.notna(value):
                chunk.append(f"{x(index):.2f},{y(float(value)):.2f}")
            elif chunk:
                chunks.append(" ".join(chunk))
                chunk = []
        if chunk:
            chunks.append(" ".join(chunk))
        segments[key] = chunks
    date_indexes = {value: index for index, value in enumerate(frame["date_key"])}
    markers = []
    for row in _records(related_signals):
        day = str(row.get("signal_date"))[:10]
        if day in date_indexes:
            markers.append(
                {
                    "x": x(date_indexes[day]),
                    "y": height - 70,
                    "kind": f"{row.get('action')}-signal",
                    "date": day,
                }
            )
    for row in _records(related_orders):
        for field, kind, marker_y in (
            ("planned_execution_date", "planned", height - 54),
            (
                "execution_date",
                "executed" if row.get("status") == "executed" else "blocked",
                height - 38,
            ),
        ):
            day = str(row.get(field))[:10]
            if day in date_indexes:
                markers.append(
                    {"x": x(date_indexes[day]), "y": marker_y, "kind": kind, "date": day}
                )
    x_count = min(6, len(frame))
    x_indexes = sorted(
        {round(item * (len(frame) - 1) / max(x_count - 1, 1)) for item in range(x_count)}
    )
    return {
        "width": width,
        "height": height,
        "symbol": symbol,
        "price_adjustment": _value(frame.iloc[0].get("price_adjustment")),
        "plot": {"left": left, "right": width - right, "top": top, "bottom": height - bottom},
        "segments": segments,
        "line_keys": line_keys,
        "markers": markers,
        "x_ticks": [{"x": x(index), "label": frame.iloc[index]["date_key"]} for index in x_indexes],
        "y_ticks": [
            {"y": y(low + span * item / 5), "label": f"{low + span * item / 5:.3f}"}
            for item in range(6)
        ],
    }


class DailySignalReportWriter:
    def __init__(self) -> None:
        self.environment = Environment(
            loader=FileSystemLoader(Path(__file__).parent / "templates"),
            autoescape=select_autoescape(["html", "xml"], default=True),
        )
        install_filters(self.environment)

    def write(
        self,
        directory: Path,
        *,
        spec: dict[str, Any],
        metrics: dict[str, Any],
        nav: pd.DataFrame,
        signals: pd.DataFrame,
        orders: pd.DataFrame,
        positions: pd.DataFrame,
        candidates: pd.DataFrame,
        diagnostics: dict[str, Any],
        account_audit: pd.DataFrame | None = None,
        reproducibility: dict[str, Any] | None = None,
        trust_summary: dict[str, Any] | None = None,
        indicator_series: pd.DataFrame | None = None,
        display_metadata: dict[str, Any] | None = None,
    ) -> dict[str, Path]:
        directory.mkdir(parents=True, exist_ok=True)
        account_audit = account_audit if account_audit is not None else pd.DataFrame()
        reproducibility = reproducibility or {}
        if trust_summary is None:
            trust_summary = daily_signal_trust_summary(
                spec=spec, metrics=metrics, nav=nav, orders=orders,
                diagnostics=diagnostics, reproducibility=reproducibility,
            )
        indicator_series = indicator_series if indicator_series is not None else pd.DataFrame()
        display_metadata = {"symbol_names": {}, **(display_metadata or {})}

        def with_columns(name: str, frame: pd.DataFrame) -> pd.DataFrame:
            expected = FRAME_COLUMNS[name]
            return frame.reindex(
                columns=[*expected, *(x for x in frame.columns if x not in expected)]
            )

        signals = with_columns("signals", signals)
        orders = with_columns("orders", orders)
        positions = with_columns("positions", positions)
        candidates = with_columns("candidates", candidates)
        paths = {
            "daily_nav_csv": directory / "daily_nav.csv",
            "signals_csv": directory / "signals.csv",
            "orders_csv": directory / "orders.csv",
            "positions_csv": directory / "positions.csv",
            "candidate_selections_csv": directory / "candidate_selections.csv",
            "metrics_json": directory / "metrics.json",
            "diagnostics_json": directory / "diagnostics.json",
            "account_audit_csv": directory / "account_audit.csv",
            "account_audit_json": directory / "account_audit.json",
            "reproducibility_json": directory / "reproducibility.json",
            "trust_summary_json": directory / "trust_summary.json",
            "indicator_series_csv": directory / "indicator_series.csv",
            "excel": directory / "daily_signal_result.xlsx",
            "html_report": directory / "report.html",
        }
        frames = {
            "daily_nav": nav,
            "signals": signals,
            "orders": orders,
            "positions": positions,
            "candidates": candidates,
            "account_audit": account_audit,
            "indicator_series": indicator_series,
        }
        for name, frame in frames.items():
            key = "candidate_selections_csv" if name == "candidates" else f"{name}_csv"
            frame.to_csv(paths[key], index=False, date_format="%Y-%m-%d")
        paths["metrics_json"].write_text(
            json.dumps(metrics, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
        )
        paths["diagnostics_json"].write_text(
            json.dumps(diagnostics, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
        )
        paths["account_audit_json"].write_text(
            json.dumps(
                diagnostics.get("account_audit", {}),
                ensure_ascii=False,
                indent=2,
                default=str,
            ),
            encoding="utf-8",
        )
        paths["reproducibility_json"].write_text(
            json.dumps(reproducibility, ensure_ascii=False, indent=2, default=str),
            encoding="utf-8",
        )
        paths["trust_summary_json"].write_text(
            json.dumps(trust_summary, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
        )
        with pd.ExcelWriter(paths["excel"], engine="openpyxl") as writer:
            pd.DataFrame([{"metric": k, "value": v} for k, v in metrics.items()]).to_excel(
                writer, sheet_name="metrics", index=False
            )
            for name, frame in frames.items():
                sheet_name = {
                    "account_audit": "每日账户核对",
                    "indicator_series": "价格指标",
                }.get(name, name)
                frame.to_excel(writer, sheet_name=sheet_name, index=False)
            pd.DataFrame(
                [{"field": key, "value": value} for key, value in reproducibility.items()]
            ).to_excel(writer, sheet_name="运行来源", index=False)
            pd.DataFrame(trust_summary.get("per_symbol") or [], columns=[
                "symbol", "buy_count", "sell_count", "buy_notional", "sell_notional", "total_cost"
            ]).to_excel(writer, sheet_name="股票成交与成本", index=False)
            for sheet in writer.book.worksheets:
                if sheet.title in {"signals", "orders", "positions", "candidates", "股票成交与成本"}:
                    headers = [cell.value for cell in sheet[1]]
                    if "symbol" in headers:
                        symbol_column = headers.index("symbol") + 1
                        name_column = sheet.max_column + 1
                        sheet.cell(1, name_column, "股票名称（运行时快照）")
                        for cells in sheet.iter_rows(min_row=2):
                            symbol = cells[symbol_column - 1].value
                            cells[name_column - 1].value = display_metadata["symbol_names"].get(str(symbol))
                sheet.freeze_panes = "A2"
                sheet.auto_filter.ref = sheet.dimensions
                for cell in sheet[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="173B5E")
                money_columns = {
                    "daily_nav": {"cash", "position_value", "portfolio_value", "total_cost"},
                    "orders": {"notional", "cost", "cash_before", "cash_after", "commission_cost", "slippage_cost", "stamp_duty_cost"},
                    "positions": {"market_value"},
                    "股票成交与成本": {"buy_notional", "sell_notional", "total_cost"},
                }.get(sheet.title, set())
                for column in sheet.iter_cols(min_row=1, max_row=1):
                    if column[0].value in money_columns:
                        for cell in sheet.iter_cols(min_col=column[0].column, max_col=column[0].column, min_row=2):
                            for value in cell:
                                if isinstance(value.value, (int, float)):
                                    value.number_format = (
                                        '0.000E+00' if value.value and abs(value.value) < 0.0000005
                                        else '#,##0.000000;[Red]-#,##0.000000;0.000000'
                                        if value.value and abs(value.value) < 0.005
                                        else '#,##0.00;[Red]-#,##0.00;0.00'
                                    )
        signal_preview = _records(signals.tail(80))
        for row in signal_preview:
            try:
                evidence = json.loads(row["rule_evidence"])
            except (TypeError, json.JSONDecodeError):
                row["rule_summary"] = "退出候选名单（candidate_drop）"
            else:
                row["rule_summary"] = "；".join(
                    f"{rule_label(item.get('rule'))}："
                    f"{'满足' if item.get('matched') is True else '不满足或数据不足'}"
                    for item in evidence
                )
        order_preview = _records(orders.tail(80))
        signal_lookup = {
            (str(row.get("signal_date")), row.get("symbol"), row.get("action")): row
            for row in signal_preview
        }
        timeline = [
            {
                "signal": signal_lookup.get(
                    (str(order.get("signal_date")), order.get("symbol"), order.get("side"))
                ),
                "order": order,
            }
            for order in order_preview
        ]
        result_view = daily_signal_result_view(
            spec=spec,
            metrics=metrics,
            nav=nav,
            orders=orders,
            positions=positions,
            diagnostics=diagnostics,
        )
        html = self.environment.get_template("daily_signal.html.j2").render(
            spec=spec,
            metrics=metrics,
            diagnostics=diagnostics,
            signals=signal_preview,
            orders=order_preview,
            timeline=timeline,
            result_view=result_view,
            nav_chart=daily_signal_chart(nav),
            price_chart=daily_signal_price_chart(
                indicator_series,
                signals,
                orders,
                symbol=(str(signals.iloc[0]["symbol"]) if not signals.empty else None),
            ),
            account_audit_summary=diagnostics.get("account_audit", {}),
            account_audit_rows=_records(account_audit.tail(100)),
            reproducibility=reproducibility,
            trust_summary=trust_summary,
            indicator_series_symbols=(
                sorted(str(value) for value in indicator_series["symbol"].dropna().unique())
                if not indicator_series.empty and "symbol" in indicator_series
                else []
            ),
            display_metadata=display_metadata,
            terms=TERM_CATALOG,
            candidates=candidates.loc[candidates.get("selected", False)].tail(40).pipe(_records)
            if not candidates.empty
            else [],
        )
        paths["html_report"].write_text(html, encoding="utf-8")
        return paths
