from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pandas as pd
from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl.styles import Alignment, Font, PatternFill

from ai_quant_assistant.reports.presentation import install_filters


class FactorBacktestReportWriter:
    def __init__(self) -> None:
        self.environment = Environment(
            loader=FileSystemLoader(Path(__file__).parent / "templates"),
            autoescape=select_autoescape(["html", "xml"], default=True),
        )
        self.environment.filters["pct"] = self._pct
        install_filters(self.environment)
        self.environment.filters["money"] = self._money

    @staticmethod
    def _pct(value: Any) -> str:
        return "-" if value is None or pd.isna(value) else f"{float(value) * 100:.2f}%"

    @staticmethod
    def _money(value: Any) -> str:
        return "-" if value is None or pd.isna(value) else f"{float(value):,.2f}"

    @staticmethod
    def _chart_points(nav: pd.DataFrame, column: str) -> str:
        if nav.empty:
            return ""
        values = pd.to_numeric(nav[column], errors="coerce").fillna(1.0)
        combined = nav[["portfolio_nav", "benchmark_nav"]].apply(pd.to_numeric, errors="coerce")
        low, high = float(combined.min().min()), float(combined.max().max())
        span = high - low or 1.0
        width, height = 960.0, 260.0
        return " ".join(
            f"{index * width / max(len(values) - 1, 1):.1f},{height - (value - low) * height / span:.1f}"
            for index, value in enumerate(values)
        )

    def write(
        self,
        directory: Path,
        *,
        spec: dict[str, Any],
        metrics: dict[str, Any],
        nav: pd.DataFrame,
        trades: pd.DataFrame,
        positions: pd.DataFrame,
        selections: pd.DataFrame,
        diagnostics: dict[str, Any],
    ) -> dict[str, Path]:
        directory.mkdir(parents=True, exist_ok=True)
        paths = {
            "daily_nav_csv": directory / "daily_nav.csv",
            "trades_csv": directory / "trades.csv",
            "positions_csv": directory / "positions.csv",
            "factor_selections_csv": directory / "factor_selections.csv",
            "metrics_json": directory / "metrics.json",
            "excel": directory / "backtest_result.xlsx",
            "html_report": directory / "report.html",
        }
        nav.to_csv(paths["daily_nav_csv"], index=False, date_format="%Y-%m-%d")
        trades.to_csv(paths["trades_csv"], index=False, date_format="%Y-%m-%d")
        positions.to_csv(paths["positions_csv"], index=False, date_format="%Y-%m-%d")
        selections.to_csv(paths["factor_selections_csv"], index=False, date_format="%Y-%m-%d")
        paths["metrics_json"].write_text(
            json.dumps(metrics, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
        )
        metric_frame = pd.DataFrame(
            [{"metric": key, "value": value} for key, value in metrics.items()]
        )
        with pd.ExcelWriter(paths["excel"], engine="openpyxl") as writer:
            for name, frame in (
                ("metrics", metric_frame),
                ("daily_nav", nav),
                ("trades", trades),
                ("positions", positions),
                ("factor_selections", selections),
            ):
                frame.to_excel(writer, sheet_name=name, index=False)
            for sheet in writer.book.worksheets:
                sheet.freeze_panes = "A2"
                if sheet.max_column and sheet.max_row:
                    sheet.auto_filter.ref = sheet.dimensions
                for cell in sheet[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="173B5E")
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                for column in sheet.columns:
                    width = min(max(len(str(cell.value or "")) for cell in column) + 2, 42)
                    sheet.column_dimensions[column[0].column_letter].width = max(10, width)
        html = self.environment.get_template("factor_backtest.html.j2").render(
            spec=spec,
            metrics=metrics,
            diagnostics=diagnostics,
            portfolio_points=self._chart_points(nav, "portfolio_nav"),
            benchmark_points=self._chart_points(nav, "benchmark_nav"),
            start_date=nav["date"].min() if not nav.empty else "",
            end_date=nav["date"].max() if not nav.empty else "",
            trades=trades.tail(30).to_dict(orient="records"),
            selections=selections.loc[selections.get("selected", False)]
            .tail(30)
            .to_dict(orient="records")
            if not selections.empty
            else [],
        )
        paths["html_report"].write_text(html, encoding="utf-8")
        return paths
