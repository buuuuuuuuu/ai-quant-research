from __future__ import annotations

import json
import sys
import time
from datetime import date
from pathlib import Path

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from test_factor_backtest import _database
from test_web_workbench import _settings

from ai_quant_assistant.api.app import create_app
from ai_quant_assistant.cli import main
from ai_quant_assistant.domain.models import FactorBacktestSpec, FactorComponent, FactorLabSpec
from ai_quant_assistant.research.factor_backtest import FactorBacktestService
from ai_quant_assistant.research.factor_lab import FactorLabService


def _spec(**changes) -> FactorLabSpec:
    options = {
        "symbols": ("600519", "000001", "300750"),
        "start_date": date(2024, 1, 2),
        "end_date": date(2024, 3, 8),
        "factors": (FactorComponent("roe"), FactorComponent("debt_asset_ratio", direction="asc")),
        "top_n_values": (1,),
        "rebalance_values": ("monthly",),
        "commission_bps_values": (5.0,),
        "initial_capital": 100_000.0,
        "quantiles": 3,
    }
    options.update(changes)
    return FactorLabSpec(**options)


def _payload(**changes) -> dict:
    payload = {
        "symbols": ["600519", "000001", "300750"],
        "start_date": "2024-01-02",
        "end_date": "2024-03-08",
        "factors": [
            {"code": "roe", "weight": 2, "direction": "desc"},
            {"code": "debt_asset_ratio", "weight": 1, "direction": "asc"},
        ],
        "top_n_values": [1],
        "rebalance_values": ["monthly"],
        "commission_bps_values": [5],
        "initial_capital": 100_000,
        "quantiles": 3,
    }
    payload.update(changes)
    return payload


@pytest.mark.parametrize(
    "changes",
    [
        {"symbols": ("600519",)},
        {"factors": ()},
        {"factors": (FactorComponent("roe"), FactorComponent("roe"))},
        {"top_n_values": (4,)},
        {"rebalance_values": ("weekly",)},
        {"commission_bps_values": (-1.0,)},
        {"quantiles": 1},
        {"training_fraction": 0.9, "validation_fraction": 0.2},
        {"missing_policy": "future"},
    ],
)
def test_factor_lab_spec_rejects_invalid_inputs(changes) -> None:
    with pytest.raises(ValueError):
        _spec(**changes)


def test_weighted_factor_snapshot_preserves_publication_cutoff_and_direction(tmp_path) -> None:
    database = _database(tmp_path)
    service = FactorBacktestService(database=database, artifact_dir=tmp_path / "artifacts")
    spec = FactorBacktestSpec(
        symbols=("600519", "000001", "300750"),
        start_date=date(2024, 1, 2),
        end_date=date(2024, 3, 8),
        factor="composite",
        factors=(
            FactorComponent("roe", weight=2),
            FactorComponent("debt_asset_ratio", direction="asc"),
        ),
        top_n=1,
    )

    before = service.factor_snapshot(spec, "2024-01-02 15:00:00")
    after = service.factor_snapshot(spec, "2024-01-03 09:00:00")

    assert set(before["symbol"]) == {"600519", "000001"}
    assert (
        json.loads(before.loc[before["symbol"] == "600519", "component_values"].iloc[0])["roe"]
        == 20.0
    )
    assert (
        json.loads(after.loc[after["symbol"] == "600519", "component_values"].iloc[0])["roe"]
        == 100.0
    )
    assert before.sort_values("factor_value", ascending=False)["symbol"].iloc[0] == "600519"


def test_composite_backtest_exports_component_lineage_and_legacy_remains_supported(
    tmp_path,
) -> None:
    database = _database(tmp_path)
    service = FactorBacktestService(database=database, artifact_dir=tmp_path / "artifacts")
    result = service.run(
        FactorBacktestSpec(
            symbols=("600519", "000001"),
            start_date=date(2024, 1, 2),
            end_date=date(2024, 3, 8),
            factor="composite",
            factors=(FactorComponent("roe"), FactorComponent("debt_asset_ratio", direction="asc")),
            top_n=1,
            rebalance="monthly",
        )
    )
    selections = pd.read_csv(result["artifacts"]["factor_selections_csv"])

    assert result["diagnostics"]["factor_combination"][1]["direction"] == "asc"
    assert selections["component_values"].notna().all()


def test_missing_factor_policy_excludes_or_renormalizes_available_components(tmp_path) -> None:
    database = _database(tmp_path)
    with database.connect() as connection:
        connection.execute(
            "DELETE FROM financial_metric WHERE security_id = 'SZSE.000001' AND metric_code = 'roe'"
        )
    service = FactorBacktestService(database=database, artifact_dir=tmp_path / "artifacts")
    base = {
        "symbols": ("600519", "000001"),
        "start_date": date(2024, 1, 2),
        "end_date": date(2024, 3, 8),
        "factor": "composite",
        "factors": (
            FactorComponent("roe", weight=2),
            FactorComponent("debt_asset_ratio", direction="asc"),
        ),
        "top_n": 1,
    }
    excluded = service.factor_snapshot(FactorBacktestSpec(**base), "2024-01-03 15:00:00")
    available = service.factor_snapshot(
        FactorBacktestSpec(**base, missing_policy="available"), "2024-01-03 15:00:00"
    )
    assert set(excluded["symbol"]) == {"600519"}
    assert set(available["symbol"]) == {"600519", "000001"}
    row = available.loc[available["symbol"] == "000001"].iloc[0]
    assert json.loads(row["component_values"])["roe"] is None
    assert pd.notna(row["published_at"])


@pytest.fixture(scope="module")
def completed_lab(tmp_path_factory):
    root = tmp_path_factory.mktemp("completed-factor-lab")
    database = _database(root)
    result = FactorLabService(database=database, artifact_dir=root / "artifacts").run(_spec())
    return root, result


def test_factor_lab_exports_ic_groups_scan_and_complete_reports(completed_lab) -> None:
    _, result = completed_lab
    assert result["status"] == "succeeded"
    assert result["summary"]["strategy_count"] == 3
    assert result["summary"]["cross_section_count"] >= 2
    assert {row["factor"] for row in result["ic_summary"]} == {
        "roe",
        "debt_asset_ratio",
        "composite_score",
    }
    assert all(Path(path).is_file() for path in result["artifacts"].values())
    workbook = load_workbook(result["artifacts"]["excel"], read_only=True)
    assert {"strategies", "segments", "ic_series", "group_returns", "walk_forward"}.issubset(
        workbook.sheetnames
    )
    assert "多因子研究与策略实验报告" in Path(result["artifacts"]["html_report"]).read_text("utf-8")


def test_factor_lab_uses_validation_only_and_keeps_future_test_separate(completed_lab) -> None:
    _, result = completed_lab
    assert result["diagnostics"]["test_period_used_for_selection"] is False
    assert [item["segment"] for item in result["validation"]] == ["train", "validation", "test"]
    train, validation, test = result["validation"]
    assert train["end_date"] < validation["start_date"] < test["start_date"]
    assert sum(row["selected_by_validation"] for row in result["strategies"]) == 1


def test_factor_lab_ic_direction_and_walk_forward_are_chronological(completed_lab) -> None:
    _, result = completed_lab
    ic = pd.read_csv(result["artifacts"]["ic_series_csv"])
    folds = pd.read_csv(result["artifacts"]["walk_forward_csv"])
    roe = ic.loc[ic["factor"] == "roe", "rank_ic"]
    debt = ic.loc[ic["factor"] == "debt_asset_ratio", "rank_ic"]
    assert roe.tolist() == debt.tolist()
    assert (folds["train_end"] < folds["test_start"]).all()


def test_factor_lab_rejects_oversized_parameter_grid_before_running(tmp_path) -> None:
    database = _database(tmp_path)
    service = FactorLabService(database=database, artifact_dir=tmp_path / "artifacts")
    with pytest.raises(ValueError, match="max_combinations"):
        service.run(_spec(top_n_values=(1, 2), max_combinations=2))


def test_factor_lab_cli_runs_local_multifactor_research(tmp_path, monkeypatch, capsys) -> None:
    _database(tmp_path)
    monkeypatch.setenv("AI_QUANT_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("AI_QUANT_DUCKDB_PATH", str(tmp_path / "data" / "catalog.duckdb"))
    monkeypatch.setenv("AI_QUANT_ARTIFACT_DIR", str(tmp_path / "artifacts"))
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "ai-quant",
            "factor-lab",
            "--symbols",
            "600519",
            "000001",
            "--start-date",
            "2024-01-02",
            "--end-date",
            "2024-03-08",
            "--factors",
            "roe:2:desc",
            "--top-n-values",
            "1",
            "--rebalance-values",
            "monthly",
        ],
    )
    main()
    result = json.loads(capsys.readouterr().out)
    assert result["summary"]["strategy_count"] == 1
    assert result["spec"]["factors"][0]["weight"] == 2


def test_factor_lab_api_background_task_history_report_and_download(tmp_path) -> None:
    _database(tmp_path)
    with TestClient(create_app(_settings(tmp_path))) as client:
        created = client.post("/api/v1/factor-lab", json=_payload())
        assert created.status_code == 202, created.text
        for _ in range(600):
            task = client.get(f"/api/v1/factor-lab/{created.json()['task_id']}").json()
            if task["status"] in {"succeeded", "failed"}:
                break
            time.sleep(0.025)
        assert task["status"] == "succeeded", task
        run_id = task["run_id"]
        detail = client.get(f"/api/v1/factor-lab-runs/{run_id}")
        history = client.get("/api/v1/factor-lab-runs")
        report = client.get(f"/reports/factor-lab/{run_id}")
        excel = client.get(f"/api/v1/factor-lab-runs/{run_id}/downloads/factor_lab.xlsx")
        tasks = client.get("/api/v1/factor-lab")
    assert detail.status_code == 200
    assert detail.json()["series"]
    assert detail.json()["group_returns"]
    assert history.json()[0]["run_id"] == run_id
    assert tasks.json()[0]["task_type"] == "factor_lab"
    assert "多因子研究与策略实验报告" in report.text
    assert excel.status_code == 200


def test_factor_lab_request_validates_factor_grid_and_secure_downloads(tmp_path) -> None:
    _database(tmp_path)
    with TestClient(create_app(_settings(tmp_path))) as client:
        unknown = client.post("/api/v1/factor-lab", json=_payload(factors=[{"code": "unknown"}]))
        excessive = client.post(
            "/api/v1/factor-lab", json=_payload(top_n_values=[1, 2], max_combinations=2)
        )
        missing = client.get("/api/v1/factor-lab-runs/not-a-run")
        traversal = client.get(
            "/api/v1/factor-lab-runs/factor-lab-20240101T000000Z-aaaaaaaa/downloads/passwords.txt"
        )
    assert unknown.status_code == 422
    assert excessive.status_code == 422
    assert missing.status_code == 404
    assert traversal.status_code == 404


def test_workbench_exposes_visual_factor_lab_and_discloses_research_bias(tmp_path) -> None:
    _database(tmp_path)
    with TestClient(create_app(_settings(tmp_path))) as client:
        home = client.get("/")
        metadata = client.get("/api")
    assert "多因子研究实验室" in home.text
    assert "验证集挑选参数" in home.text
    assert "幸存者偏差" in home.text
    assert "lab-top-n" in home.text
    assert "股票池与行业数据" in home.text
    assert "lab-universe-mode" in home.text
    assert "可信历史指数成分" in home.text
    assert metadata.json()["factor_lab"] == "/api/v1/factor-lab"
    assert metadata.json()["universe"] == "/api/v1/universe/overview"
