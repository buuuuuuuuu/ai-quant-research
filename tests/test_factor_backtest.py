from __future__ import annotations

import json
import sys
from datetime import date
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from ai_quant_assistant.cli import main
from ai_quant_assistant.domain.models import FactorBacktestSpec
from ai_quant_assistant.research.factor_backtest import FactorBacktestService
from ai_quant_assistant.storage.database import LocalDatabase
from ai_quant_assistant.storage.layout import DataLayout


def _database(tmp_path: Path) -> LocalDatabase:
    data_dir = tmp_path / "data"
    database = LocalDatabase(path=data_dir / "catalog.duckdb", layout=DataLayout.from_root(data_dir))
    database.initialize()
    trading_days = pd.bdate_range("2024-01-02", "2024-04-10")
    securities = (
        ("SSE.600519", "600519", "SSE", "2001-08-27", 100.0, 0.30),
        ("SZSE.000001", "000001", "SZSE", "1991-04-03", 50.0, 0.10),
        ("SZSE.300750", "300750", "SZSE", "2024-03-01", 80.0, 0.20),
    )
    with database.connect() as connection:
        for security_id, symbol, exchange, listed, base, slope in securities:
            connection.execute(
                """
                INSERT INTO security_master (
                    security_id, symbol, exchange, name, list_date, source, source_batch_id
                ) VALUES (?, ?, ?, ?, ?, 'fixture-backtest', 'fixture-batch')
                """,
                [security_id, symbol, exchange, f"测试公司{symbol}", listed],
            )
            bars = []
            for index, day in enumerate(trading_days):
                if day.date() < date.fromisoformat(listed):
                    continue
                close = base + index * slope
                for adjustment in ("none", "qfq"):
                    bars.append(
                        [
                            security_id,
                            day.date(),
                            adjustment,
                            close,
                            close + 0.20,
                            close - 0.20,
                            close,
                            1_000.0,
                            close * 1_000,
                        ]
                    )
            connection.executemany(
                """
                INSERT INTO daily_bar (
                    security_id, trade_date, adjustment, open, high, low, close,
                    volume, amount, source_batch_id, quality_status
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'fixture-batch', 'published')
                """,
                bars,
            )

        connection.execute(
            """
            INSERT INTO benchmark_master (benchmark_id, symbol, name, source, source_batch_id)
            VALUES ('INDEX.csi000300', 'csi000300', '沪深300', 'fixture', 'fixture-batch')
            """
        )
        connection.executemany(
            """
            INSERT INTO benchmark_daily_bar (
                benchmark_id, trade_date, close, source, source_batch_id
            ) VALUES ('INDEX.csi000300', ?, ?, 'fixture', 'fixture-batch')
            """,
            [[day.date(), 1_000.0 + index * 0.5] for index, day in enumerate(trading_days)],
        )

        reports = (
            ("maotai-q3", "SSE.600519", "2023-09-30", "q3", "2023-12-15 18:00:00", 20.0),
            ("maotai-revision", "SSE.600519", "2023-09-30", "q3", "2024-01-02 18:00:00", 100.0),
            ("pingan-q3", "SZSE.000001", "2023-09-30", "q3", "2023-12-16 18:00:00", 10.0),
            (
                "pingan-annual",
                "SZSE.000001",
                "2023-12-31",
                "annual",
                "2024-02-01 18:00:00",
                150.0,
            ),
            (
                "catl-annual",
                "SZSE.300750",
                "2023-12-31",
                "annual",
                "2024-02-28 18:00:00",
                5.0,
            ),
        )
        for report_id, security_id, period, report_type, published, roe in reports:
            connection.execute(
                """
                INSERT INTO financial_report (
                    report_id, security_id, report_period_end, report_type, published_at,
                    availability_source, source, source_batch_id, raw_sha256
                ) VALUES (?, ?, ?, ?, ?, 'provider_notice', 'fixture', 'fixture-batch', 'sha')
                """,
                [report_id, security_id, period, report_type, published],
            )
            for factor, value in (("roe", roe), ("debt_asset_ratio", 200.0 - roe)):
                connection.execute(
                    """
                    INSERT INTO financial_metric (
                        report_id, security_id, report_period_end, published_at,
                        metric_code, metric_value, unit, source_batch_id
                    ) VALUES (?, ?, ?, ?, ?, ?, 'PERCENT', 'fixture-batch')
                    """,
                    [report_id, security_id, period, published, factor, value],
                )
        from managed_test_data import admit_fixture

        connection.execute(
            """INSERT INTO trading_calendar(exchange,trade_date,is_open,source,source_batch_id)
               SELECT exchange, CAST(day AS DATE), TRUE, 'fixture-backtest', 'fixture-batch'
               FROM generate_series(DATE '2024-01-02',DATE '2024-04-10',INTERVAL 1 DAY) days(day)
               CROSS JOIN (VALUES ('SSE'),('SZSE')) exchanges(exchange)
               WHERE isodow(day) <= 5 ON CONFLICT DO NOTHING"""
        )
        admit_fixture(connection, [item[1] for item in securities])
    return database


def _spec(**changes) -> FactorBacktestSpec:
    options = {
        "symbols": ("600519", "000001", "300750"),
        "start_date": date(2024, 1, 2),
        "end_date": date(2024, 3, 8),
        "factor": "roe",
        "top_n": 1,
        "rebalance": "monthly",
        "initial_capital": 100_000.0,
        "commission_bps": 5.0,
        "slippage_bps": 5.0,
        "stamp_duty_bps": 5.0,
    }
    options.update(changes)
    return FactorBacktestSpec(**options)


def _service(tmp_path: Path) -> tuple[FactorBacktestService, LocalDatabase]:
    database = _database(tmp_path)
    return FactorBacktestService(database=database, artifact_dir=tmp_path / "artifacts"), database


@pytest.mark.parametrize(
    "changes",
    (
        {"symbols": ()},
        {"symbols": ("bad",)},
        {"top_n": 0},
        {"top_n": 4},
        {"rebalance": "weekly"},
        {"factor_direction": "sideways"},
        {"initial_capital": 0.0},
        {"commission_bps": -1.0},
        {"lot_size": 0},
    ),
)
def test_backtest_spec_rejects_invalid_inputs(changes) -> None:
    with pytest.raises(ValueError):
        _spec(**changes)


def test_factor_snapshot_excludes_after_close_publication_and_respects_revisions(tmp_path) -> None:
    database = _database(tmp_path)

    before = database.read_factor_snapshot(("600519",), "roe", "2024-01-02 15:00:00")
    after = database.read_factor_snapshot(("600519",), "roe", "2024-01-03 09:00:00")
    before_listing = database.read_factor_snapshot(("300750",), "roe", "2024-02-29 15:00:00")
    after_listing = database.read_factor_snapshot(("300750",), "roe", "2024-03-01 15:00:00")

    assert before["factor_value"].tolist() == [20.0]
    assert after["factor_value"].tolist() == [100.0]
    assert before_listing.empty
    assert after_listing["factor_value"].tolist() == [5.0]


def test_monthly_backtest_uses_next_day_execution_and_no_future_financials(tmp_path) -> None:
    service, _ = _service(tmp_path)

    result = service.run(_spec())
    selections = pd.read_csv(result["artifacts"]["factor_selections_csv"], dtype={"symbol": str})
    trades = pd.read_csv(result["artifacts"]["trades_csv"], dtype={"symbol": str})
    selected = selections.loc[selections["selected"]]

    assert result["status"] == "succeeded"
    assert selected.loc[selected["signal_date"] == "2024-01-02", "symbol"].tolist() == ["600519"]
    assert selected.loc[selected["signal_date"] == "2024-02-01", "symbol"].tolist() == ["600519"]
    assert selected.loc[selected["signal_date"] == "2024-03-01", "symbol"].tolist() == ["000001"]
    assert trades.iloc[0]["date"] == "2024-01-03"
    assert trades.iloc[0]["signal_date"] == "2024-01-02"
    march = trades.loc[trades["date"] == "2024-03-04"]
    assert march["side"].tolist() == ["sell", "buy"]
    assert result["diagnostics"]["point_in_time_rule"].endswith("15:00")


def test_backtest_charges_costs_and_enforces_round_lot_buys(tmp_path) -> None:
    service, _ = _service(tmp_path)

    with_costs = service.run(_spec())
    without_costs = service.run(
        _spec(commission_bps=0.0, slippage_bps=0.0, stamp_duty_bps=0.0)
    )
    trades = pd.read_csv(with_costs["artifacts"]["trades_csv"])
    buys = trades.loc[(trades["status"] == "executed") & (trades["side"] == "buy")]

    assert with_costs["metrics"]["total_transaction_cost"] > 0
    assert without_costs["metrics"]["total_transaction_cost"] == 0
    assert with_costs["metrics"]["final_value"] < without_costs["metrics"]["final_value"]
    assert all(int(quantity) % 100 == 0 for quantity in buys["shares"])


def test_one_price_limit_up_blocks_purchase(tmp_path) -> None:
    service, database = _service(tmp_path)
    with database.connect() as connection:
        connection.execute(
            """
            UPDATE daily_bar SET open = 110, high = 110, low = 110, close = 110
            WHERE security_id = 'SSE.600519' AND trade_date = DATE '2024-01-03'
            """
        )

    result = service.run(_spec(end_date=date(2024, 1, 10)))
    trades = pd.read_csv(result["artifacts"]["trades_csv"])

    assert trades.iloc[0]["status"] == "blocked"
    assert trades.iloc[0]["reason"] == "one_price_limit_up"
    assert result["diagnostics"]["blocked_order_reasons"]["one_price_limit_up"] == 1


def test_suspension_blocks_purchase_without_fake_fill(tmp_path) -> None:
    service, database = _service(tmp_path)
    with database.connect() as connection:
        connection.execute(
            """
            UPDATE daily_bar SET volume = 0
            WHERE security_id = 'SSE.600519' AND trade_date = DATE '2024-01-03'
            """
        )

    result = service.run(_spec(end_date=date(2024, 1, 10)))
    trades = pd.read_csv(result["artifacts"]["trades_csv"])

    assert trades.iloc[0]["status"] == "blocked"
    assert trades.iloc[0]["reason"] == "suspended_or_missing_bar"
    assert result["metrics"]["trade_count"] == 0


def test_one_price_limit_down_blocks_sale_and_preserves_existing_position(tmp_path) -> None:
    service, database = _service(tmp_path)
    with database.connect() as connection:
        previous = connection.execute(
            """
            SELECT close FROM daily_bar
            WHERE security_id = 'SSE.600519'
              AND trade_date = DATE '2024-03-01'
              AND adjustment = 'none'
            """
        ).fetchone()[0]
        limit_price = float(previous) * 0.9
        connection.execute(
            """
            UPDATE daily_bar SET open = ?, high = ?, low = ?, close = ?
            WHERE security_id = 'SSE.600519' AND trade_date = DATE '2024-03-04'
            """,
            [limit_price, limit_price, limit_price, limit_price],
        )

    result = service.run(_spec())
    trades = pd.read_csv(result["artifacts"]["trades_csv"], dtype={"symbol": str})
    blocked = trades.loc[(trades["symbol"] == "600519") & (trades["status"] == "blocked")]
    positions = pd.read_csv(result["artifacts"]["positions_csv"], dtype={"symbol": str})

    assert blocked["reason"].tolist() == ["one_price_limit_down"]
    assert "600519" in positions.loc[positions["date"] == "2024-03-04", "symbol"].tolist()
    assert result["diagnostics"]["blocked_order_reasons"]["one_price_limit_down"] == 1


def test_ascending_factor_sort_and_quarterly_rebalance(tmp_path) -> None:
    service, _ = _service(tmp_path)

    result = service.run(
        _spec(factor="debt_asset_ratio", factor_direction="asc", rebalance="quarterly")
    )
    selections = pd.read_csv(result["artifacts"]["factor_selections_csv"], dtype={"symbol": str})

    assert selections.loc[selections["selected"], "symbol"].tolist() == ["600519"]
    assert result["metrics"]["rebalance_count"] == 1


def test_backtest_exports_complete_research_artifacts(tmp_path) -> None:
    service, _ = _service(tmp_path)

    result = service.run(_spec())
    paths = result["artifacts"]
    workbook = load_workbook(paths["excel"], read_only=True)
    report = Path(paths["html_report"]).read_text(encoding="utf-8")
    manifest = json.loads(Path(paths["run_manifest"]).read_text(encoding="utf-8"))
    nav = pd.read_csv(paths["daily_nav_csv"])

    assert all(Path(path).is_file() for path in paths.values())
    assert workbook.sheetnames == [
        "metrics",
        "daily_nav",
        "trades",
        "positions",
        "factor_selections",
    ]
    assert "财务因子组合回测" in report
    assert "下一基准交易日开盘" in report
    assert manifest["spec"]["factor"] == "roe"
    assert nav["portfolio_nav"].iloc[0] == 1.0
    assert nav["benchmark_nav"].iloc[0] == 1.0
    assert result["metrics"]["max_drawdown"] >= 0


def test_backtest_requires_published_local_benchmark(tmp_path) -> None:
    service, database = _service(tmp_path)
    with database.connect() as connection:
        connection.execute("DELETE FROM benchmark_daily_bar")

    with pytest.raises(ValueError, match="benchmark"):
        service.run(_spec())


def test_cli_runs_factor_backtest_entirely_from_local_database(
    tmp_path, monkeypatch, capsys
) -> None:
    _database(tmp_path)
    monkeypatch.setenv("AI_QUANT_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("AI_QUANT_ARTIFACT_DIR", str(tmp_path / "artifacts"))
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "ai-quant",
            "backtest",
            "--source",
            "local",
            "--symbols",
            "600519",
            "000001",
            "--start-date",
            "2024-01-02",
            "--end-date",
            "2024-03-08",
            "--factor",
            "roe",
            "--top-n",
            "1",
            "--rebalance",
            "monthly",
        ],
    )

    main()
    payload = json.loads(capsys.readouterr().out)

    assert payload["status"] == "succeeded"
    assert payload["spec"]["factor"] == "roe"
    assert payload["metrics"]["trade_count"] >= 1
    assert Path(payload["artifacts"]["html_report"]).is_file()
