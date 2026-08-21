from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl.styles import Alignment, Font, PatternFill


class BatchReportWriter:
    def __init__(self) -> None:
        self.environment = Environment(
            loader=FileSystemLoader(Path(__file__).parent / "templates"),
            autoescape=select_autoescape(["html", "xml"]),
        )
        self.environment.filters["pct"] = self._pct

    @staticmethod
    def _pct(value: Any) -> str:
        if value is None or pd.isna(value):
            return "-"
        return f"{float(value) * 100:.2f}%"

    def write(
        self,
        directory: Path,
        *,
        spec: dict[str, Any],
        pooled_summary: pd.DataFrame,
        stock_summary: pd.DataFrame,
        events: pd.DataFrame,
        failures: pd.DataFrame,
    ) -> dict[str, Path]:
        directory.mkdir(parents=True, exist_ok=True)
        paths = {
            "pooled_summary_csv": directory / "pooled_summary.csv",
            "stock_summary_csv": directory / "stock_summary.csv",
            "events_csv": directory / "events.csv",
            "failures_csv": directory / "failures.csv",
            "excel": directory / "batch_result.xlsx",
            "html_report": directory / "batch_report.html",
        }
        pooled_summary.to_csv(paths["pooled_summary_csv"], index=False)
        stock_summary.to_csv(paths["stock_summary_csv"], index=False)
        events.to_csv(paths["events_csv"], index=False, date_format="%Y-%m-%d")
        failures.to_csv(paths["failures_csv"], index=False)
        with pd.ExcelWriter(paths["excel"], engine="openpyxl") as writer:
            pooled_summary.to_excel(writer, sheet_name="pooled_summary", index=False)
            stock_summary.to_excel(writer, sheet_name="stock_summary", index=False)
            events.to_excel(writer, sheet_name="events", index=False)
            failures.to_excel(writer, sheet_name="failures", index=False)
            for sheet in writer.book.worksheets:
                sheet.freeze_panes = "A2"
                if sheet.max_column and sheet.max_row:
                    sheet.auto_filter.ref = sheet.dimensions
                for cell in sheet[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="173B5E")
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                for column in sheet.columns:
                    max_length = min(max(len(str(cell.value or "")) for cell in column) + 2, 42)
                    sheet.column_dimensions[column[0].column_letter].width = max(10, max_length)
        html = self.environment.get_template("batch_event_study.html.j2").render(
            spec=spec,
            pooled_summary=pooled_summary.to_dict(orient="records"),
            stock_summary=stock_summary.to_dict(orient="records"),
            failures=failures.to_dict(orient="records"),
            event_count=len(events),
            success_count=stock_summary["symbol"].nunique() if "symbol" in stock_summary else 0,
        )
        paths["html_report"].write_text(html, encoding="utf-8")
        return paths
